"""Module 1 + 2a: validate an uploaded resume and extract its raw text.

Files are read from memory (bytes), so nothing is written to disk.
"""
from io import BytesIO

from docx import Document
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".pdf", ".docx"}
MAX_FILE_SIZE_MB = 5


class ResumeError(ValueError):
    """Raised when a resume cannot be accepted or read."""


def validate_file(filename: str, size_bytes: int) -> str:
    """Check extension and size; return the lowercase extension."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise ResumeError(f"Unsupported file type '{ext or 'none'}'. Upload a PDF or DOCX file.")
    if size_bytes == 0:
        raise ResumeError("The file is empty.")
    if size_bytes > MAX_FILE_SIZE_MB * 1024 * 1024:
        raise ResumeError(f"File is larger than {MAX_FILE_SIZE_MB} MB.")
    return ext


def extract_pdf_text(data: bytes) -> str:
    """Extract text from every page of a PDF."""
    try:
        reader = PdfReader(BytesIO(data))
    except Exception as exc:  # pypdf raises several error types for broken files
        raise ResumeError(f"Could not open the PDF: {exc}") from exc
    if reader.is_encrypted:
        raise ResumeError("The PDF is password protected.")
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def extract_docx_text(data: bytes) -> str:
    """Extract text from every paragraph and table cell of a DOCX file."""
    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise ResumeError(f"Could not open the DOCX file: {exc}") from exc
    lines = [p.text for p in document.paragraphs]
    # Many resume templates put skills or contact details in tables.
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    """Validate the file and return its raw text."""
    ext = validate_file(filename, len(data))
    text = extract_pdf_text(data) if ext == ".pdf" else extract_docx_text(data)
    if len(text.strip()) < 50:
        raise ResumeError(
            "Very little text was found. The resume may be a scanned image; "
            "upload a text-based PDF or DOCX."
        )
    return text
