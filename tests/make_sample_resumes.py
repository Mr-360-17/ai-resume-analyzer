"""Generate the anonymised sample resumes in data/sample_resumes/.

All candidates are fictional. Personal details are replaced with placeholders,
as the brief requires. Run:  python tests/make_sample_resumes.py
"""
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).resolve().parent.parent / "data" / "sample_resumes"

RESUMES = {
    "resume_A_data_analyst": {
        "format": "pdf",
        "title": "Candidate A - Aspiring Data Analyst",
        "sections": {
            "Contact": ["Email: [removed] | Phone: [removed] | City: [removed]"],
            "Education": ["B.Sc. Statistics, 2025 - CGPA 8.4"],
            "Skills": ["Python, Pandas, NumPy, SQL (MySQL), Excel (pivot tables, VLOOKUP), "
                       "Power BI, Tableau, Statistics, Data Visualization"],
            "Projects": [
                "Sales Performance Dashboard - cleaned 50k rows of retail data with Pandas and "
                "built an interactive Power BI dashboard with DAX measures for monthly KPIs.",
                "Customer Churn EDA - exploratory data analysis in Python with Matplotlib and "
                "Seaborn; used hypothesis testing to compare churn across plans.",
                "SQL Practice Project - wrote JOIN and GROUP BY queries on a MySQL database "
                "to answer business questions; exported results to Excel.",
            ],
            "Experience": ["Data Analyst Intern, 3 months - automated a weekly Excel report "
                           "and handled data cleaning of missing values."],
        },
    },
    "resume_B_ml_engineer": {
        "format": "docx",
        "title": "Candidate B - Machine Learning Engineer",
        "sections": {
            "Contact": ["Email: [removed] | Phone: [removed]"],
            "Education": ["B.Tech Computer Science, 2024"],
            "Technical Skills": ["Python, scikit-learn, XGBoost, Pandas, NumPy, SQL, FastAPI, "
                                 "Docker, MLflow, Git, GitHub, AWS"],
            "Projects": [
                "Loan Default Predictor - trained a gradient boosting model with XGBoost, "
                "tracked experiments in MLflow and served it through a FastAPI REST API.",
                "House Price Regression - feature engineering, cross-validation and a "
                "scikit-learn pipeline; containerised with Docker and deployed on AWS EC2.",
            ],
            "Work Experience": ["Machine Learning Intern, 6 months - built classification "
                                "models with scikit-learn and set up CI/CD with GitHub Actions."],
        },
    },
    "resume_C_nlp_engineer": {
        "format": "pdf",
        "title": "Candidate C - NLP Engineer",
        "sections": {
            "Contact": ["Email: [removed] | LinkedIn: [removed]"],
            "Education": ["M.Tech Artificial Intelligence, 2025"],
            "Skills": ["Python, NLP, spaCy, NLTK, Hugging Face Transformers, PyTorch, "
                       "TF-IDF, word embeddings, scikit-learn"],
            "Projects": [
                "Movie Review Sentiment Analysis - fine-tuned BERT with Hugging Face and PyTorch "
                "for text classification; compared against a TF-IDF + logistic regression baseline.",
                "Resume Entity Extractor - named entity recognition with spaCy to pull "
                "skills and organisations from resumes.",
                "News Topic Classifier - NLTK preprocessing and word2vec embeddings.",
            ],
            "Experience": ["Research Assistant, 1 year - built NLP pipelines for text "
                           "classification in Python."],
        },
    },
    "resume_D_cv_engineer": {
        "format": "docx",
        "title": "Candidate D - Computer Vision Engineer",
        "sections": {
            "Contact": ["Email: [removed] | Phone: [removed]"],
            "Education": ["B.E. Electronics and Communication, 2024"],
            "Skills": ["Python, OpenCV, PyTorch, TensorFlow, CNN, YOLO, Deep Learning, NumPy"],
            "Projects": [
                "Helmet Detection - trained YOLOv8 for object detection on traffic camera "
                "footage; bounding boxes evaluated with mAP.",
                "Leaf Disease Classifier - transfer learning with a ResNet convolutional neural "
                "network in PyTorch; image processing with OpenCV.",
                "Road Segmentation - U-Net image segmentation in TensorFlow.",
            ],
            "Experience": ["Computer Vision Intern, 4 months - deep learning models for "
                           "defect detection on a factory line."],
        },
    },
    "resume_E_web_developer": {
        "format": "pdf",
        "title": "Candidate E - Frontend Web Developer",
        "sections": {
            "Contact": ["Email: [removed] | Portfolio: [removed]"],
            "Education": ["BCA, 2025"],
            "Skills": ["HTML5, CSS3, JavaScript (ES6), React, TypeScript, Tailwind, Git, GitHub, "
                       "REST APIs, Node.js"],
            "Projects": [
                "Recipe Finder - React app that calls a public REST API; styled with Tailwind CSS.",
                "Portfolio Website - responsive HTML and CSS site hosted on GitHub Pages.",
                "Task Board - TypeScript + React frontend with a Node.js and Express backend.",
            ],
            "Experience": ["Web Development Intern, 3 months - built React components and "
                           "fixed CSS layout bugs."],
        },
    },
}


def write_pdf(path: Path, title: str, sections: dict[str, list[str]]) -> None:
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"])]
    for heading, lines in sections.items():
        story.append(Paragraph(heading, styles["Heading2"]))
        for line in lines:
            story.append(Paragraph(line, styles["BodyText"]))
            story.append(Spacer(1, 4))
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


def write_docx(path: Path, title: str, sections: dict[str, list[str]]) -> None:
    doc = Document()
    doc.add_heading(title, level=0)
    for heading, lines in sections.items():
        doc.add_heading(heading, level=1)
        for line in lines:
            doc.add_paragraph(line, style="List Bullet" if heading.endswith("Projects") else None)
    doc.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, spec in RESUMES.items():
        path = OUT / f"{name}.{spec['format']}"
        writer = write_pdf if spec["format"] == "pdf" else write_docx
        writer(path, spec["title"], spec["sections"])
        print("wrote", path.name)


if __name__ == "__main__":
    main()
